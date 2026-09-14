#!/usr/bin/env python3
"""
De projectenlijst van H-Architects met fototelling omzetten naar JSON.

Eenmalig werk, geen onderdeel van de draaiende app: net als het VEKA-register en
het architectenregister wordt het bronbestand met de hand ververst. De bron is de
werklijst die op 14 september 2026 uit Dropbox is opgebouwd:

    ~/Claude/marketing/seo/firmas/H-Architects/H-Architects-projectenlijst-fotos-2026-09-14.xlsx

De toelichting bij die telling (methode, wat "foto's aanwezig" wel en niet
betekent, welke mappen niet leesbaar waren) staat in het LEESMIJ-bestand ernaast
en is samengevat onderaan /h-architects/projecten.

    python3 scripts/h-architects-projecten.py [pad-naar-xlsx]

Vereist openpyxl (`python3 -m pip install openpyxl`, of een venv).

Let op: de mapnamen bevatten klant- en medewerkersnamen. Dit bestand hoort
daarom uitsluitend in deze privé-repo, achter Authentik -- net als het
VEKA-register. Niet exporteren, niet op een website zetten.
"""

import json, os, sys

try:
    import openpyxl
except ImportError:  # pragma: no cover - alleen een nette melding
    sys.exit("openpyxl ontbreekt: python3 -m pip install openpyxl")

HIER = os.path.dirname(os.path.abspath(__file__))
UIT = os.path.join(HIER, "..", "data-bronnen", "h-architects-projecten-2026-09-14.json")
STANDAARD_BRON = os.path.expanduser(
    "~/Claude/marketing/seo/firmas/H-Architects/"
    "H-Architects-projectenlijst-fotos-2026-09-14.xlsx"
)

# Datum van de telling zelf (Dropbox-doorloop) en van de projectlijst eronder.
TELLING = "2026-09-14"
PROJECTLIJST = "2026-09-09"

# De zes projectmappen waarvan Dropbox de inhoud niet wil geven. Ze staan in de
# lijst met foto's = "onbekend", uitdrukkelijk niet als "nee".
ONLEESBARE_MAPPEN = [
    {"nummer": "1840", "project": "Brusselstraat 63, Antwerpen"},
    {"nummer": "1842", "project": "Statielei 18, Mortsel"},
    {"nummer": "1931", "project": "Brouwerijstraat 48, Lokeren"},
    {"nummer": "1960", "project": "Grote Kauwenberg 13, Antwerpen"},
    {"nummer": "2198", "project": "Grimaldilaan 15, Stabroek"},
    {"nummer": "2419", "project": "Wolfshaegen 118, Neerijse (voorstudie, geannuleerd)"},
]

# Vier gedeelde mappen geven met dit Dropbox-account nog altijd not_found.
# Projecten die uitsluitend daar zitten, staan dus niet in de lijst.
GEDEELDE_MAPPEN = [
    "/Work All/01. H-A WORK/1 H-A Opzegging Architectuurovereenkomst",
    "/Work All/01. H-A WORK/1 H-A Geannuleerd",
    "/Work All/01. H-A WORK/H-A Bouwteam",
    "/Work All/01. H-A WORK/0 H-A Standaard projects/H-A Mehdi (oude projecten voor de facturatie)",
]

KOLOMMEN = [
    "nummer", "adres", "postcode", "gemeente", "provincie", "jaar", "type",
    "status", "fotos", "aantal", "recentste", "oplevering", "_link", "mapnaam",
]


def tekst(v):
    if v is None:
        return ""
    return str(v).strip()


def getal(v):
    if v is None or v == "":
        return None
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def main():
    bron = sys.argv[1] if len(sys.argv) > 1 else STANDAARD_BRON
    wb = openpyxl.load_workbook(bron)
    ws = wb["Projecten"]

    kop = [tekst(c.value) for c in ws[1]]
    if len(kop) != len(KOLOMMEN):
        sys.exit(f"onverwachte kolommen in het tabblad Projecten: {kop}")

    projecten = []
    for rij in ws.iter_rows(min_row=2):
        cel = dict(zip(KOLOMMEN, rij))
        # De link zit als hyperlink op de cel ("map openen"), niet in de tekst.
        link = cel["_link"].hyperlink.target if cel["_link"].hyperlink else ""
        if link and not link.startswith("https://www.dropbox.com/"):
            sys.exit(f"onverwachte link: {link}")

        ruw_type = tekst(cel["type"].value)
        projecten.append({
            "nummer": tekst(cel["nummer"].value),
            "adres": tekst(cel["adres"].value),
            "postcode": getal(cel["postcode"].value),
            "gemeente": tekst(cel["gemeente"].value),
            "provincie": tekst(cel["provincie"].value),
            "jaar": getal(cel["jaar"].value),
            # Eén project kan meerdere typen dragen ("regularisatie, light").
            # Als lijst bewaren zodat het filter op één type kan matchen.
            "typen": [t.strip() for t in ruw_type.split(",") if t.strip()],
            "status": tekst(cel["status"].value),
            "fotos": tekst(cel["fotos"].value),          # ja / nee / onbekend
            "aantal": getal(cel["aantal"].value) or 0,
            "recentste": tekst(cel["recentste"].value) or None,  # ISO-datum
            "oplevering": getal(cel["oplevering"].value) or 0,
            "link": link,
            "mapnaam": tekst(cel["mapnaam"].value),
        })

    uit = {
        "bron": os.path.basename(bron),
        "telling": TELLING,
        "projectlijst": PROJECTLIJST,
        "onleesbareMappen": ONLEESBARE_MAPPEN,
        "gedeeldeMappen": GEDEELDE_MAPPEN,
        "projecten": projecten,
    }

    pad = os.path.normpath(UIT)
    with open(pad, "w", encoding="utf-8") as f:
        json.dump(uit, f, ensure_ascii=False, indent=1)
        f.write("\n")

    met = sum(1 for p in projecten if p["fotos"] == "ja")
    zonder = sum(1 for p in projecten if p["fotos"] == "nee")
    onbekend = sum(1 for p in projecten if p["fotos"] == "onbekend")
    oplev = sum(1 for p in projecten if p["oplevering"] > 0)
    recent = sum(1 for p in projecten if (p["recentste"] or "") >= "2024-01-01")
    zonderlink = sum(1 for p in projecten if not p["link"])
    print(f"{pad}: {len(projecten)} projecten")
    print(f"  met foto {met} · zonder {zonder} · onbekend {onbekend}")
    print(f"  opleveringsfoto's {oplev} · beeld 2024+ {recent} · zonder link {zonderlink}")


if __name__ == "__main__":
    main()
